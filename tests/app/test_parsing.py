import io
import pytest
from app.parsing import parse_file, ParseError, UnsupportedFormat


def test_parse_txt_and_md():
    assert parse_file("a.txt", "你好\n世界".encode()) == "你好\n世界"
    assert "标题" in parse_file("b.md", "# 标题".encode())


def test_parse_docx_roundtrip():
    import docx
    d = docx.Document()
    d.add_paragraph("第一段内容")
    d.add_paragraph("第二段内容")
    buf = io.BytesIO(); d.save(buf)
    text = parse_file("x.docx", buf.getvalue())
    assert "第一段内容" in text and "第二段内容" in text


def test_parse_docx_includes_table_cells():
    """回归：docx 只取 paragraphs 会漏掉表格里的全部文字（那些在 tables 里）。

    简历、报表、对照表大量用表格，漏了它们上传后知识库检索不到——这是静默的数据丢失。
    """
    import docx
    d = docx.Document()
    d.add_paragraph("正文开头")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "姓名"; t.cell(0, 1).text = "年龄"
    t.cell(1, 0).text = "张三"; t.cell(1, 1).text = "28"
    buf = io.BytesIO(); d.save(buf)
    text = parse_file("resume.docx", buf.getvalue())
    for cell in ("姓名", "年龄", "张三", "28"):
        assert cell in text, f"表格单元格「{cell}」丢失"


def test_parse_docx_preserves_paragraph_table_order():
    """段落与表格在 python-docx 里是两条平行列表，简单拼接会打乱先后。
    须按文档真实顺序输出：正文 → 表格 → 正文。"""
    import docx
    d = docx.Document()
    d.add_paragraph("前言段落")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "表格内容"
    d.add_paragraph("结尾段落")
    buf = io.BytesIO(); d.save(buf)
    text = parse_file("x.docx", buf.getvalue())
    assert text.index("前言段落") < text.index("表格内容") < text.index("结尾段落")


def test_parse_docx_includes_nested_table():
    """单元格里可以再嵌一张表（复杂报表）。cell.text 只读单元格段落、不读嵌套表，
    须递归下钻，否则内层整表静默丢失。"""
    import docx
    d = docx.Document()
    outer = d.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "外层单元格"
    inner = cell.add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "内层嵌套表格文字"
    buf = io.BytesIO(); d.save(buf)
    text = parse_file("report.docx", buf.getvalue())
    assert "外层单元格" in text and "内层嵌套表格文字" in text


def test_parse_docx_includes_content_control():
    """内容控件 w:sdt（表单/模板的下拉、日期框）在 body 级是独立元素，把段落包在
    w:sdtContent 里；平铺遍历会整段跳过，须钻进 sdtContent。"""
    import docx
    from docx.oxml.shared import OxmlElement, qn
    d = docx.Document()
    sdt = OxmlElement("w:sdt"); content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p"); r = OxmlElement("w:r"); t = OxmlElement("w:t")
    t.text = "内容控件里的文字"
    r.append(t); p.append(r); content.append(p); sdt.append(content)
    d.element.body.append(sdt)
    buf = io.BytesIO(); d.save(buf)
    assert "内容控件里的文字" in parse_file("form.docx", buf.getvalue())


def test_parse_docx_omits_header_footer_noise():
    """页眉/页脚在独立 XML part，多为页码/水印/装饰——刻意不抽，免得给知识库塞噪声。
    这条同时守住「递归下钻别顺手把它们也引进来」。"""
    import docx
    d = docx.Document()
    d.add_paragraph("正文内容")
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "页眉水印"
    sec.footer.paragraphs[0].text = "第 1 页"
    buf = io.BytesIO(); d.save(buf)
    text = parse_file("x.docx", buf.getvalue())
    assert "正文内容" in text
    assert "页眉水印" not in text and "第 1 页" not in text


def test_unsupported_format_raises():
    with pytest.raises(UnsupportedFormat):
        parse_file("x.pptx", b"data")


def test_none_filename_raises_unsupported():
    # UploadFile.filename 可能为 None，应干净抛 UnsupportedFormat 而非 TypeError
    with pytest.raises(UnsupportedFormat):
        parse_file(None, b"data")


def test_corrupt_pdf_raises_parse_error():
    with pytest.raises(ParseError):
        parse_file("broken.pdf", b"not a real pdf")


def test_corrupt_docx_raises_parse_error():
    with pytest.raises(ParseError):
        parse_file("broken.docx", b"not a real docx")
