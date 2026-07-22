# app/parsing.py
from __future__ import annotations

import io


class UnsupportedFormat(Exception):
    pass


class ParseError(Exception):
    """文件格式受支持但内容损坏/无法解析。"""


def parse_file(filename: str, data: bytes) -> str:
    name = filename or ""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    if ext == "pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:  # pypdf 对损坏文件抛多种异常
            raise ParseError(f"PDF 解析失败：{e}") from e
    if ext == "docx":
        import docx
        try:
            document = docx.Document(io.BytesIO(data))
            return _docx_text(document)
        except Exception as e:  # python-docx 对非 zip/损坏文件抛 PackageNotFoundError 等
            raise ParseError(f"docx 解析失败：{e}") from e
    raise UnsupportedFormat(ext or name)


def _docx_text(document) -> str:
    """按文档原始顺序抽出段落与表格文字。

    只取 document.paragraphs 会漏掉表格里的全部文字（那些在 document.tables 里，是一条
    独立的列表）——简历、报表、对照表这类 docx 大量用表格，漏了它们上传后知识库检索不到。
    而 paragraphs 与 tables 是两个平行列表，简单拼接会打乱正文与表格的先后；故遍历
    body 的 XML 子元素按真实顺序取，段落原样、表格逐行用制表符连列。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            parts.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, document).rows:
                parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)
